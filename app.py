from flask import Flask, render_template, request, redirect, url_for, send_file
from werkzeug.utils import secure_filename
import os
import pdfplumber
from openpyxl import Workbook
import re
from docx import Document

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def extract_info_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        text = ""
        for page in pdf.pages:
            text += page.extract_text()

    email_regex = r'[\w\.-]+@[\w\.-]+'
    phone_regex = r'[\+\(]?[1-9][0-9 .\-\(\)]{8,}[0-9]'

    emails = re.findall(email_regex, text)
    phones = re.findall(phone_regex, text)

    return emails, phones, text

def extract_info_from_docx(docx_path):
    doc = Document(docx_path)
    text = "\n".join([paragraph.text for paragraph in doc.paragraphs])

    email_regex = r'[\w\.-]+@[\w\.-]+'
    phone_regex = r'[\+\(]?[1-9][0-9 .\-\(\)]{8,}[0-9]'

    emails = re.findall(email_regex, text)
    phones = re.findall(phone_regex, text)

    return emails, phones, text

def generate_excel(data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["File", "Email", "Phone", "Text"])

    for filename, emails, phones, text in data:
        email_str = ", ".join(set(emails)) if emails else ""
        phone_str = ", ".join(set(phones)) if phones else ""
        text_lines = text.split("\n")

        if not text_lines:
            continue

        # Append the first line of text with email and phone info
        ws.append([filename, email_str, phone_str, text_lines[0]])

        # Append subsequent lines of text without email and phone info
        for line in text_lines[1:]:
            ws.append(["", "", "", line])

    # Autofit column widths
    for column in ws.columns:
        max_length = 0
        column = list(column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column[0].column_letter].width = adjusted_width

    wb.save(output_path)
    print("Excel file generated successfully.")

@app.route('/', methods=['GET', 'POST'])
def upload_file():
    if request.method == 'POST':
        file = request.files['file']
        if file.filename == '':
            return render_template('index.html', error='No file selected')
        if file:
            filename = secure_filename(file.filename)
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            output_file = os.path.join(app.config['UPLOAD_FOLDER'], 'cv_info.xlsx')
            
            if filename.endswith('.pdf'):
                emails, phones, text = extract_info_from_pdf(filepath)
            elif filename.endswith('.docx'):
                emails, phones, text = extract_info_from_docx(filepath)
            else:
                return render_template('index.html', error='Unsupported file format')
            
            generate_excel([(filename, emails, phones, text)], output_file)
            
            return redirect(url_for('download', filename='cv_info.xlsx'))
    return render_template('index.html')

@app.route('/download/<filename>')
def download(filename):
    return send_file(os.path.join(app.config['UPLOAD_FOLDER'], filename), as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
